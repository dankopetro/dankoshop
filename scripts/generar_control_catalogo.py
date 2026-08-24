#!/usr/bin/env python3
"""
Genera Excel de control cruzando Catálogo WhatsApp con Maestro actual.

Columnas:
  A: SKU (Maestro o SINxxx)
  B: Artículo (Maestro o catálogo)
  C: Categoría (Maestro)
  D: Descripción (Maestro)
  E: Precio Mayorista (catálogo - control)
  F: Inventario (Maestro)
  G: Envío Grande (Maestro)
  H: Canal de Venta (Maestro)
  I: Precio Medusa (mayorista × 1.0606 × 1.072) → sincroniza a Medusa
  J: Precio Público Contado (Precio Medusa × 1.15) → control
  K: Precio Venta Crédito (Precio Medusa × 1.25) → control

Uso:
  python3 scripts/generar_control_catalogo.py
"""

import re
from pathlib import Path
from difflib import SequenceMatcher
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE = Path(__file__).resolve().parent.parent
CATALOGO_PATH = Path("/home/claudio/Descargas/Mayorista/24Agosto/Catalogo_WhatsApp_Completo.xlsx")
MAESTRO_PATH = BASE / "excel" / "Productos_Maestro.xlsx"
OUTPUT_PATH = BASE / "excel" / "Control_Catalogo_24_Agosto.xlsx"

FACTOR_PROV = 1.0606
FACTOR_SYNC = 1.072
FACTOR_CONTADO = 1.15
FACTOR_CREDITO = 1.25


def parse_price(val):
    if val is None:
        return None
    s = str(val).replace("ARS", "").replace("\xa0", "").replace(",", "").replace("$", "").strip()
    try:
        return float(s)
    except Exception:
        return None


def normalize(s):
    return re.sub(r"[^a-z0-9 ]", "", s.lower()).strip()


def load_catalogo():
    wb = openpyxl.load_workbook(CATALOGO_PATH, data_only=True)
    ws = wb["Catálogo"]
    items = []
    for r in range(2, ws.max_row + 1):
        art = str(ws.cell(r, 2).value or "").strip()
        precio = parse_price(ws.cell(r, 3).value)
        if art and precio is not None:
            items.append({"row": r, "articulo": art, "precio_catalogo": precio})
    return items


def load_maestro():
    wb = openpyxl.load_workbook(MAESTRO_PATH, data_only=True)
    ws = wb.active
    items = []
    for r in range(2, ws.max_row + 1):
        sku = str(ws.cell(r, 1).value or "").strip()
        if not sku:
            continue
        items.append({
            "sku": sku,
            "articulo": str(ws.cell(r, 2).value or "").strip(),
            "categoria": str(ws.cell(r, 3).value or "").strip(),
            "descripcion": str(ws.cell(r, 4).value or "").strip(),
            "precio_mayorista": ws.cell(r, 6).value,
            "envio_grande": ws.cell(r, 7).value or "",
            "inventario": ws.cell(r, 8).value,
            "canales": str(ws.cell(r, 9).value or "").strip(),
        })
    return items


def find_last_sin(maestro):
    max_num = 0
    for m in maestro:
        if m["sku"].startswith("SIN"):
            try:
                num = int(m["sku"][3:])
                max_num = max(max_num, num)
            except ValueError:
                pass
    return max_num


def best_match(catalog_name, maestro):
    cn = normalize(catalog_name)
    best_score = 0
    best_m = None
    for m in maestro:
        mn = normalize(m["articulo"])
        if cn in mn or mn in cn:
            score = 0.9
        else:
            score = SequenceMatcher(None, cn, mn).ratio()
        if score > best_score:
            best_score = score
            best_m = m
    return best_m, best_score


def main():
    print("Cargando catálogo...")
    catalogo = load_catalogo()
    print(f"  {len(catalogo)} productos con precio")

    print("Cargando maestro...")
    maestro = load_maestro()
    print(f"  {len(maestro)} productos")

    last_sin = find_last_sin(maestro)
    print(f"  Último SKU SIN: SIN{last_sin:03d}")

    # Header colors
    HEADER_COLORS = {
        "SKU": "D35230",
        "Artículo": "E8A838",
        "Categoría": "4CAF50",
        "Descripción": "2196F3",
        "Precio Mayorista": "FF9800",
        "Inventario": "795548",
        "Envío Grande": "00BCD4",
        "Canal de Venta": "607D8B",
        "Precio Medusa": "9C27B0",
        "Precio Público Contado": "4CAF50",
        "Precio Venta Crédito": "F44336",
    }
    HEADERS = [
        "SKU", "Artículo", "Categoría", "Descripción",
        "Precio Mayorista", "Inventario", "Envío Grande", "Canal de Venta",
        "Precio Medusa", "Precio Público Contado", "Precio Venta Crédito",
    ]

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center")
    data_center = Alignment(horizontal="center", vertical="center")
    data_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Control Catálogo"

    # Headers
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = center
        color = HEADER_COLORS.get(h, "333333")
        cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        cell.border = thin_border

    # Match and write
    new_sin = last_sin
    matched = 0
    new_products = 0
    no_price = 0
    used_skus = set()

    for idx, cat_item in enumerate(catalogo, 2):
        art = cat_item["articulo"]
        precio_catalogo = cat_item["precio_catalogo"]

        m, score = best_match(art, maestro)

        if score >= 0.5 and m and m["sku"] not in used_skus:
            sku = m["sku"]
            cat_art = m["articulo"]
            categoria = m["categoria"]
            descripcion = m["descripcion"]
            inv = m["inventario"]
            envio = m["envio_grande"]
            canales = m["canales"]
            used_skus.add(sku)
            matched += 1
        else:
            new_sin += 1
            sku = f"SIN{new_sin:03d}"
            cat_art = art
            categoria = ""
            descripcion = ""
            inv = None
            envio = ""
            canales = ""
            new_products += 1

        precio_proveedor = precio_catalogo
        precio_mayorista_calc = int(round(precio_proveedor * FACTOR_PROV))
        precio_medusa = int(round(precio_proveedor * FACTOR_PROV * FACTOR_SYNC))
        precio_contado = int(round(precio_proveedor * FACTOR_PROV * FACTOR_SYNC * FACTOR_CONTADO))
        precio_credito = int(round(precio_proveedor * FACTOR_PROV * FACTOR_SYNC * FACTOR_CREDITO))

        row_data = [
            sku, cat_art, categoria, descripcion,
            precio_catalogo, inv if inv is not None else "", envio, canales,
            precio_medusa, precio_contado, precio_credito,
        ]

        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=idx, column=col, value=val)
            cell.border = thin_border
            if col in (1, 6, 7):
                cell.alignment = data_center
            elif col in (5, 9, 10, 11):
                cell.alignment = data_center
                cell.number_format = '#,##0'
            else:
                cell.alignment = data_left

    # Column widths
    widths = [14, 55, 28, 45, 18, 12, 14, 22, 18, 22, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{len(catalogo) + 1}"

    out.save(OUTPUT_PATH)
    print(f"\nResultado:")
    print(f"  Matcheados con Maestro: {matched}")
    print(f"  Productos nuevos (SIN):  {new_products}")
    print(f"  Total:                   {matched + new_products}")
    print(f"\nGuardado: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
