#!/usr/bin/env python3
"""
Sync products from Medusa v2 backend to excel/Productos_Maestro.xlsx.
Usage:
  MEDUSA_BACKEND_URL=https://... python3 sync_medusa_to_excel.py

Reads all products from Medusa, extracts metadata and prices,
and writes them to the Excel master file.
"""

import json
import os
import sys
import requests
from pathlib import Path

# Config
MEDUSA_URL = os.getenv("MEDUSA_BACKEND_URL", "http://localhost:9100")
ADMIN_EMAIL = os.getenv("MEDUSA_ADMIN_EMAIL", "admin@dankoshop.com")
ADMIN_PASSWORD = os.getenv("MEDUSA_ADMIN_PASSWORD", "supersecret")
EXCEL_PATH = Path(__file__).parent.parent / "excel" / "Productos_Maestro.xlsx"

MEDUSA_API_TOKEN = None
HEADERS = {"Content-Type": "application/json"}


def get_auth_token():
    global MEDUSA_API_TOKEN
    try:
        resp = requests.post(
            f"{MEDUSA_URL}/auth/user/emailpass",
            json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
            headers={"Content-Type": "application/json"},
            timeout=10
        )
        if resp.status_code == 200:
            MEDUSA_API_TOKEN = resp.json().get("token")
            return MEDUSA_API_TOKEN
    except Exception as e:
        print(f"Login failed: {e}")
    return None


def medusa_request(method, endpoint, data=None, params=None):
    token = get_auth_token()
    if not token:
        raise Exception("No auth token")
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    return requests.request(method, f"{MEDUSA_URL}{endpoint}", headers=headers, json=data, params=params, timeout=30)


def fetch_all_products():
    products = []
    offset = 0
    limit = 100
    while True:
        resp = medusa_request("GET", "/admin/products", params={"limit": limit, "offset": offset, "fields": "*categories,*variants,*images"})
        if resp.status_code != 200:
            print(f"Error fetching products: {resp.status_code}")
            break
        batch = resp.json().get("products", [])
        products.extend(batch)
        if len(batch) < limit:
            break
        offset += limit
    return products


def get_category_name(product):
    cats = product.get("categories", [])
    if cats:
        cat = cats[0]
        if isinstance(cat, dict):
            return cat.get("name") or cat.get("category", {}).get("name", "General")
        return str(cat)
    return "General"


def sync_to_excel():
    import openpyxl

    products = fetch_all_products()
    print(f"Fetched {len(products)} products from Medusa")

    if EXCEL_PATH.exists():
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active

        existing_skus = {}
        for row in ws.iter_rows(min_row=2, values_only=False):
            if row[0].value:
                existing_skus[str(row[0].value).strip()] = row
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Productos"
        headers = ["SKU", "Nombre", "Categoría", "Precio Lista", "Precio Efectivo",
                    "Precio Transferencia", "Precio Mayorista", "Cuotas", "Valor Cuota",
                    "Descripción", "Imágenes"]
        ws.append(headers)
        existing_skus = {}

    updated = 0
    created = 0
    for p in products:
        for v in p.get("variants", []):
            sku = v.get("sku")
            if not sku:
                continue

            meta = p.get("metadata") or {}
            lista = meta.get("precio_lista")
            efectivo = meta.get("precio_efectivo")
            transferencia = meta.get("precio_transferencia")
            mayorista = meta.get("precio_mayorista")
            c3 = meta.get("cuota_valor_3")

            images_list = p.get("images", [])
            img_urls = " | ".join(img.get("url", "") for img in images_list if img.get("url"))
            if not img_urls and p.get("thumbnail"):
                img_urls = p["thumbnail"]

            category = get_category_name(p)
            description = p.get("description") or ""

            if sku in existing_skus:
                row = existing_skus[sku]
                row[1].value = p.get("title", "")
                row[2].value = category
                row[3].value = lista
                row[4].value = efectivo
                row[5].value = transferencia
                row[6].value = mayorista
                row[7].value = 3
                row[8].value = c3
                row[9].value = description
                row[10].value = img_urls
                updated += 1
            else:
                new_row = [sku, p.get("title", ""), category, lista, efectivo,
                           transferencia, mayorista, 3, c3, description, img_urls]
                ws.append(new_row)
                existing_skus[sku] = ws.max_row
                created += 1

            break

    wb.save(EXCEL_PATH)
    print(f"Done: {updated} updated, {created} created in {EXCEL_PATH}")


if __name__ == "__main__":
    sync_to_excel()
