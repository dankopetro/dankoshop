#!/usr/bin/env python3
"""Finaliza un Control corregido manualmente y genera Control + Sync.

Uso:
  python3 scripts/finalizar_control_corregido.py \
    '/ruta/Control_21_08_2026.xlsx'

Las correcciones manuales de SKU, Artículo, Precio fuente o Compra Mayorista
son respetadas. Si Compra Mayorista está cargada, es la autoridad final;
Venta Bruta se recalcula como Compra * 1.072.
"""

import argparse
import os
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
import requests
from openpyxl.styles import Font, PatternFill
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
from excel_style import style_workbook

BASE = Path(__file__).resolve().parent.parent
MEDUSA_URL = os.getenv("MEDUSA_BACKEND_URL", "https://dankoshop-api-production.up.railway.app")
MEDUSA_EMAIL = os.getenv("MEDUSA_ADMIN_EMAIL", "admin@dankoshop.com")
MEDUSA_PASSWORD = os.getenv("MEDUSA_ADMIN_PASSWORD", "supersecret")


def norm(value):
    value = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9]+", " ", value.upper())).strip()


def number(value):
    try:
        return round(float(str(value).replace("$", "").replace(",", "").strip()))
    except (TypeError, ValueError):
        return None


def envio_flag(value):
    return "TRUE" if str(value or "").strip().upper() in {"TRUE", "1", "SI", "SÍ", "YES", "X"} else "FALSE"


def medusa_index(backup_path=None):
    """Obtiene los precios actuales de Medusa para replicar su criterio de sync."""
    if backup_path:
        try:
            data = __import__("json").loads(Path(backup_path).read_text(encoding="utf-8"))
            result = {}
            for product in data.get("products", []):
                metadata = product.get("metadata") or {}
                for variant in product.get("variants", []):
                    sku = str(variant.get("sku") or "").strip().upper()
                    if sku:
                        result[sku] = metadata
            return result, True
        except (OSError, ValueError, TypeError):
            return {}, False
    try:
        auth = requests.post(f"{MEDUSA_URL}/auth/user/emailpass",
                             json={"email": MEDUSA_EMAIL, "password": MEDUSA_PASSWORD}, timeout=15)
        if auth.status_code != 200:
            return {}, False
        token = auth.json().get("token")
        headers = {"Authorization": f"Bearer {token}"}
        result = {}
        # El endpoint de búsqueda devuelve variantes y metadata, igual que el sincronizador.
        page = 0
        while True:
            response = requests.get(f"{MEDUSA_URL}/admin/products",
                                    params={"limit": 100, "offset": page * 100},
                                    headers=headers, timeout=20)
            if response.status_code != 200:
                return {}, False
            products = response.json().get("products", [])
            for product in products:
                metadata = product.get("metadata") or {}
                for variant in product.get("variants", []):
                    sku = str(variant.get("sku") or "").strip().upper()
                    if sku:
                        result[sku] = metadata
            if len(products) < 100:
                break
            page += 1
        return result, True
    except requests.RequestException:
        return {}, False


def master_indexes(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
    cat_col = next((c for c, h in enumerate(headers, 1) if h in ("categoría", "categoria")), None)
    inv_col = next((c for c, h in enumerate(headers, 1) if h == "inventario"), None)
    sc_col = next((c for c, h in enumerate(headers, 1) if "canales" in h), None)
    by_sku, by_name = {}, defaultdict(list)
    for r in range(2, ws.max_row + 1):
        sku = str(ws.cell(r, 1).value or "").strip()
        name = str(ws.cell(r, 2).value or "").strip()
        if not name:
            continue
        item = {
            "sku": sku, "name": name, "price": ws.cell(r, 6).value, "envio": ws.cell(r, 7).value,
            "categoria": str(ws.cell(r, cat_col).value or "").strip() if cat_col else "",
            "inventario": ws.cell(r, inv_col).value if inv_col else None,
            "canales": str(ws.cell(r, sc_col).value or "").strip() if sc_col else "",
        }
        by_sku[sku.upper()] = item
        by_name[norm(name)].append(item)
    return by_sku, by_name


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("control", type=Path)
    ap.add_argument("--maestro", type=Path, default=BASE / "excel" / "Productos_Maestro.xlsx")
    ap.add_argument("--out-dir", type=Path, default=None)
    ap.add_argument("--medusa-backup", type=Path, default=None,
                    help="Backup JSON de Medusa para clasificar cambios sin consultar el servidor")
    args = ap.parse_args()
    out_dir = args.out_dir or args.control.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = args.control.stem.replace("_CORREGIDO", "")
    control_out = out_dir / f"{stem}_CORREGIDO.xlsx"
    sync_out = out_dir / f"Nuevos_{stem.replace('Control_', '')}_Sync.xlsx"

    by_sku, by_name = master_indexes(args.maestro)
    medusa_prices, medusa_ok = medusa_index(args.medusa_backup)
    source_wb = openpyxl.load_workbook(args.control, data_only=True)
    ws = source_wb["Control"] if "Control" in source_wb.sheetnames else source_wb.active
    headers = {str(cell.value).strip(): i + 1 for i, cell in enumerate(ws[1]) if cell.value}

    def col(name, fallback):
        return headers.get(name, fallback)

    c_fecha, c_hora = col("Fecha", 1), col("Hora", 2)
    c_sku, c_other = col("SKU", 3), col("Otros SKUs", 4)
    c_name, c_img = col("Artículo", 5), col("Imagen", 6)
    c_source, c_compra = col("Precio fuente", 7), col("Compra Mayorista", 9)
    # Los Controles antiguos no tenían esta columna: en ese caso se toma siempre del Maestro.
    c_envio = headers.get("Envío Grande")

    rows, sync_rows, review, seen = [], [], [], set()
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, c_name).value or "").strip()
        if not name:
            continue
        sku = str(ws.cell(r, c_sku).value or "").strip()
        match = by_sku.get(sku.upper()) if sku else None
        status = "SKU OK" if match else ""
        if not match:
            candidates = by_name.get(norm(name), [])
            if candidates:
                match = candidates[0]
                sku = match["sku"]
                status = "SKU COMPLETADO POR NOMBRE"
        if not match:
            status = "REVISAR SKU"

        source = number(ws.cell(r, c_source).value)
        compra = number(ws.cell(r, c_compra).value)
        if compra is None and source is not None:
            compra = round(source * 1.0606)
        venta = round(compra * 1.072) if compra is not None else ""
        envio_input = ws.cell(r, c_envio).value if c_envio else None
        envio = envio_flag(envio_input) if str(envio_input or "").strip() else envio_flag(match.get("envio") if match else None)
        if compra is None:
            status = "REVISAR PRECIO"
        if sku.upper() in seen and sku:
            status = "DUPLICADO"
        else:
            seen.add(sku.upper())

        old = match.get("price") if match else ""
        diff = round((compra - old) / old * 100, 2) if compra is not None and number(old) else ""
        if compra is None:
            accion = "REVISAR PRECIO"
        elif not match:
            accion = "CREAR / REVISAR PRODUCTO"
        elif medusa_ok and sku.upper() not in medusa_prices:
            accion = "CREAR / REVISAR PRODUCTO"
        elif medusa_ok:
            current_metadata = medusa_prices.get(sku.upper())
            if current_metadata is None:
                accion = "CREAR / REVISAR PRODUCTO"
            else:
                nuevo = {
                    "precio_lista": round(compra * 1.25),
                    "precio_efectivo": round(compra * 1.15),
                    "precio_mayorista": round(compra),
                    "cuota_valor_3": round(round(compra * 1.25) / 3),
                    "cuota_valor_6": round(round(compra * 1.25) / 6),
                }
                # Replica también el cálculo de cuotas a 12 del sincronizador.
                tasa = 0.12
                factor = ((1 + tasa) ** 12 - 1) / (tasa * (1 + tasa) ** 12)
                nuevo["cuota_valor_12"] = round(round(compra * 1.25) / factor)
                if envio == "TRUE":
                    nuevo["envio_grande"] = True
                cambios = any(current_metadata.get(k) != v for k, v in nuevo.items())
                if envio != "TRUE" and "envio_grande" in current_metadata:
                    cambios = True
                accion = "ACTUALIZAR PRECIO" if cambios else "SIN CAMBIOS"
        elif number(old) is not None and compra != number(old):
            accion = "ACTUALIZAR PRECIO"
        else:
            accion = "SIN CAMBIOS"
        row = [
            ws.cell(r, c_fecha).value, ws.cell(r, c_hora).value, sku,
            ws.cell(r, c_other).value or "", name, ws.cell(r, c_img).value or "",
            source or "", "", compra or "", venta, old, diff, status,
        ]
        row.insert(10, envio)
        rows.append(row)
        review.append([r, sku, name, compra or "", status, accion])
        if status != "DUPLICADO" and sku and compra is not None:
            sync_rows.append([sku, name, compra, ws.cell(r, c_fecha).value or "", ws.cell(r, c_img).value or "", venta, envio,
                              match.get("categoria", "") if match else "",
                              match.get("inventario") if match else None,
                              match.get("canales", "") if match else ""])

    out = openpyxl.Workbook()
    cw = out.active
    cw.title = "Control"
    control_headers = ["Fecha", "Hora", "SKU", "Otros SKUs", "Artículo", "Imagen", "Precio fuente", "Fuente", "Compra Mayorista", "Venta Bruta", "Envío Grande", "Maestro F", "Diff %", "Estado"]
    cw.append(control_headers)
    for cell in cw[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="34495E")
    for row in rows:
        cw.append(row)
    cw.freeze_panes = "A2"
    cw.auto_filter.ref = f"A1:N{cw.max_row}"
    for i, width in enumerate([12, 8, 24, 18, 58, 35, 16, 18, 20, 18, 14, 16, 10, 28], 1):
        cw.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    sw = out.create_sheet("Nuevos Sync")
    sw.append(["SKU", "Artículo", "Precio Compra Mayorista", "Fecha Actualización", "Imágenes JPG", "Precio Venta Bruto", "Envío Grande", "Categoría", "Inventario", "Canales de Venta"])
    for row in sync_rows:
        sw.append(row)
    sw.freeze_panes = "A2"
    sw.auto_filter.ref = f"A1:J{sw.max_row}"
    for i, width in enumerate([24, 62, 22, 18, 35, 18, 14, 28, 12, 22], 1):
        sw.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    rw = out.create_sheet("Revisión")
    accion_header = "Acción en Medusa" if medusa_ok else "Acción estimada (Maestro; Medusa no disponible)"
    rw.append(["Fila original", "SKU final", "Artículo", "Compra final", "Estado", accion_header])
    for row in review:
        rw.append(row)
    rw.freeze_panes = "A2"
    rw.auto_filter.ref = f"A1:F{rw.max_row}"
    style_workbook(out)
    out.save(control_out)

    # El Sync se guarda como archivo independiente, listo para Medusa.
    sync_wb = openpyxl.Workbook()
    sync_ws = sync_wb.active
    sync_ws.title = "Nuevos Sync"
    sync_ws.append(["SKU", "Artículo", "Precio Compra Mayorista", "Fecha Actualización", "Imágenes JPG", "Precio Venta Bruto", "Envío Grande", "Categoría", "Inventario", "Canales de Venta"])
    for row in sync_rows:
        sync_ws.append(row)
    sync_ws.freeze_panes = "A2"
    sync_ws.auto_filter.ref = f"A1:J{sync_ws.max_row}"
    for i, width in enumerate([24, 62, 22, 18, 35, 18, 14, 28, 12, 22], 1):
        sync_ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    style_workbook(sync_wb)
    sync_wb.save(sync_out)
    print(f"Control corregido: {control_out}")
    print(f"Sync definitivo: {sync_out}")
    print(f"Productos: {len(sync_rows)} | Revisados: {len(review)}")


if __name__ == "__main__":
    main()
