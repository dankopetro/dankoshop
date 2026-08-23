#!/usr/bin/env python3
"""Valida y finaliza un Sync corregido manualmente contra Productos_Maestro.xlsx.

Uso:
  python3 scripts/finalizar_sync_corregido.py excel/Nuevos_19_08_2026_Sync.xlsx

El precio de Compra Mayorista corregido manualmente es la fuente final;
Venta Bruta se recalcula como Compra * 1.072.
"""

import argparse
import re
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
from excel_style import style_workbook

BASE = Path(__file__).resolve().parent.parent


def norm(value):
    value = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9]+", " ", value.upper())).strip()


def envio_flag(value):
    return "TRUE" if str(value or "").strip().upper() in {"TRUE", "1", "SI", "SÍ", "YES", "X"} else "FALSE"


def load_master(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
    cat_col = next((c for c, h in enumerate(headers, 1) if h in ("categoría", "categoria")), None)
    inv_col = next((c for c, h in enumerate(headers, 1) if h == "inventario"), None)
    sc_col = next((c for c, h in enumerate(headers, 1) if "canales" in h), None)
    by_sku, by_name = {}, {}
    for r in range(2, ws.max_row + 1):
        sku, name = str(ws.cell(r, 1).value or "").strip(), str(ws.cell(r, 2).value or "").strip()
        if not name:
            continue
        item = {
            "sku": sku, "name": name, "row": r, "envio": ws.cell(r, 7).value,
            "categoria": str(ws.cell(r, cat_col).value or "").strip() if cat_col else "",
            "inventario": ws.cell(r, inv_col).value if inv_col else None,
            "canales": str(ws.cell(r, sc_col).value or "").strip() if sc_col else "",
        }
        if sku:
            by_sku[sku.upper()] = item
        by_name.setdefault(norm(name), []).append(item)
    return by_sku, by_name


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("sync", type=Path)
    ap.add_argument("--maestro", type=Path, default=BASE / "excel" / "Productos_Maestro.xlsx")
    ap.add_argument("--output", type=Path, default=None)
    args = ap.parse_args()
    if args.output:
        # Acepta tanto --output carpeta como --output ruta/archivo.xlsx.
        if args.output.exists() and args.output.is_dir():
            output = args.output / (args.sync.stem + "_DEFINITIVO.xlsx")
        elif args.output.suffix.lower() != ".xlsx":
            args.output.mkdir(parents=True, exist_ok=True)
            output = args.output / (args.sync.stem + "_DEFINITIVO.xlsx")
        else:
            output = args.output
            output.parent.mkdir(parents=True, exist_ok=True)
    else:
        output = args.sync.with_name(args.sync.stem + "_DEFINITIVO.xlsx")

    by_sku, by_name = load_master(args.maestro)
    wb = openpyxl.load_workbook(args.sync, data_only=True)
    ws = wb["Nuevos Sync"] if "Nuevos Sync" in wb.sheetnames else wb.active

    rows, review, seen = [], [], set()
    for r in range(2, ws.max_row + 1):
        sku = str(ws.cell(r, 1).value or "").strip()
        name = str(ws.cell(r, 2).value or "").strip()
        compra = ws.cell(r, 3).value
        fecha = ws.cell(r, 4).value or ""
        imagen = ws.cell(r, 5).value or ""
        envio_input = ws.cell(r, 7).value
        match = by_sku.get(sku.upper()) if sku else None
        status = "SKU OK" if match else ""
        if not match:
            candidates = by_name.get(norm(name), [])
            if candidates:
                match = candidates[0]
                sku = match["sku"]
                status = "SKU COMPLETADO POR NOMBRE"
        if not match:
            status = "REVISAR SKU"
        envio = envio_flag(envio_input) if str(envio_input or "").strip() else envio_flag(match.get("envio") if match else None)
        try:
            compra_num = round(float(compra))
            venta = round(compra_num * 1.072)
        except (TypeError, ValueError):
            compra_num, venta = "", ""
            status = "REVISAR PRECIO"
        if sku and sku.upper() in seen:
            review.append([r, sku, name, compra_num, "DUPLICADO — no se incluyó"])
            continue
        if sku:
            seen.add(sku.upper())
        rows.append([sku, name, compra_num, fecha, imagen, venta, envio,
                      match.get("categoria", "") if match else "",
                      match.get("inventario") if match else None,
                      match.get("canales", "") if match else ""])
        review.append([r, sku, name, compra_num, status])

    out = openpyxl.Workbook()
    sws = out.active
    sws.title = "Nuevos Sync"
    sws.append(["SKU", "Artículo", "Precio Compra Mayorista", "Fecha Actualización", "Imágenes JPG", "Precio Venta Bruto", "Envío Grande", "Categoría", "Inventario", "Canales de Venta"])
    for row in sws[1]:
        row.font = Font(bold=True, color="FFFFFF")
        row.fill = PatternFill("solid", fgColor="34495E")
    for row in rows:
        sws.append(row)
    sws.freeze_panes = "A2"
    sws.auto_filter.ref = f"A1:J{sws.max_row}"
    for col, width in enumerate([22, 62, 22, 18, 35, 18, 14, 28, 12, 22], 1):
        sws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    rev = out.create_sheet("Revisión")
    rev.append(["Fila original", "SKU final", "Artículo", "Compra corregida", "Estado"])
    for row in rev[1]:
        row.font = Font(bold=True, color="FFFFFF")
        row.fill = PatternFill("solid", fgColor="34495E")
    for row in review:
        rev.append(row)
    rev.freeze_panes = "A2"
    rev.auto_filter.ref = f"A1:E{rev.max_row}"
    for col, width in enumerate([16, 22, 62, 20, 30], 1):
        rev.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    style_workbook(out)
    out.save(output)
    print(f"Definitivo: {output}")
    print(f"Productos incluidos: {len(rows)} | Revisados: {len(review)}")


if __name__ == "__main__":
    main()
