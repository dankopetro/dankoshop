#!/usr/bin/env python3
"""Genera planillas Control + Sync desde un chat exportado de WhatsApp.

Uso:
  python3 scripts/generar_planillas_desde_chat.py /ruta/chat.txt
  python3 scripts/generar_planillas_desde_chat.py /ruta/chat.txt \
      --images-dir /ruta/Imagenes --date 2026-08-19

Regla de precio:
  transferencia (incluido "Mayorista por transferencia") > mayorista > efectivo
  > precio único/x1. Al precio elegido se aplica 1.0606 y luego 1.072.
"""

import argparse
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).resolve().parent.parent
FACTOR_COMPRA = 1.0606
FACTOR_VENTA = 1.072

sys.path.insert(0, str(Path(__file__).resolve().parent))
from procesar_chat_14_15 import (  # noqa: E402
    extract_skus,
    find_title,
    parse_prices,
    parse_timestamp,
    split_by_timestamp,
    strip_timestamp,
)
from excel_style import style_workbook  # noqa: E402


def norm_text(value):
    value = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Za-z0-9]+", " ", value.upper())).strip()


def numeric(value):
    return isinstance(value, (int, float)) and value > 0


def envio_flag(value):
    """Normaliza la marca de envío grande del Maestro a TRUE/FALSE."""
    return "TRUE" if str(value or "").strip().upper() in {"TRUE", "1", "SI", "SÍ", "YES", "X"} else "FALSE"


def choose_source(prices):
    """Choose the message price according to the business rule."""
    choices = [
        ("transferencia", prices.get("mayorista_transfer")),
        ("transferencia", prices.get("precio_trans_minorista")),
        ("mayorista", prices.get("mayorista_efectivo")),
        ("efectivo", prices.get("precio_unidad")),
        ("unico/x1", prices.get("precio_lista")),
    ]
    for source, value in choices:
        if numeric(value):
            return value, source
    return None, ""


def load_maestro(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
    cat_col = next((c for c, h in enumerate(headers, 1) if h in ("categoría", "categoria")), None)
    inv_col = next((c for c, h in enumerate(headers, 1) if h == "inventario"), None)
    sc_col = next((c for c, h in enumerate(headers, 1) if "canales" in h), None)
    by_sku, by_name = {}, defaultdict(list)
    for row in range(2, ws.max_row + 1):
        sku = str(ws.cell(row, 1).value or "").strip()
        name = str(ws.cell(row, 2).value or "").strip()
        if not name:
            continue
        item = {"row": row, "sku": sku, "name": name, "price": ws.cell(row, 6).value,
                "envio": ws.cell(row, 7).value,
                "categoria": str(ws.cell(row, cat_col).value or "").strip() if cat_col else "",
                "inventario": ws.cell(row, inv_col).value if inv_col else None,
                "canales": str(ws.cell(row, sc_col).value or "").strip() if sc_col else ""}
        if sku:
            by_sku[sku.upper()] = item
        by_name[norm_text(name)].append(item)
    return by_sku, by_name


def image_map(images_dir):
    if not images_dir or not images_dir.exists():
        return {}
    out = defaultdict(list)
    for path in images_dir.iterdir():
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        match = re.search(r"(\d{4}-\d{2}-\d{2} at \d{2}\.\d{2}\.\d{2})", path.name)
        if match:
            out[match.group(1)].append(path.name)
    return out


def parse_chat(chat_path, images_dir=None):
    lines = chat_path.read_text(encoding="utf-8").splitlines(True)
    parsed, images = [], image_map(images_dir)
    used_images = defaultdict(int)
    for message in split_by_timestamp(lines):
        timestamp = message.get("timestamp")
        raw_lines = message.get("lines", [])
        body = [strip_timestamp(raw_lines[0])] + [line.strip() for line in raw_lines[1:] if line.strip()]
        skus = extract_skus(raw_lines)[1]
        prices = parse_prices(raw_lines)
        source_price, source_kind = choose_source(prices)
        title = find_title(body)
        if not title or not source_price:
            continue
        folder = timestamp.strftime("%Y-%m-%d at %H.%M.%S") if timestamp else ""
        image = ""
        if folder in images and used_images[folder] < len(images[folder]):
            image = images[folder][used_images[folder]]
            used_images[folder] += 1
        parsed.append({
            "date": timestamp.strftime("%d/%m/%Y") if timestamp else "",
            "time": timestamp.strftime("%H:%M") if timestamp else "",
            "sku": skus[0] if skus else "",
            "other_skus": ", ".join(skus[1:]),
            "name": title,
            "image": image,
            "source": source_price,
            "source_kind": source_kind,
            "compra": round(source_price * FACTOR_COMPRA),
            "venta": round(source_price * FACTOR_COMPRA * FACTOR_VENTA),
        })
    return parsed


def compare_and_dedupe(rows, maestro_path):
    by_sku, by_name = load_maestro(maestro_path)
    result, seen = [], set()
    for item in rows:
        key = item["sku"].upper() or norm_text(item["name"])
        if key in seen:
            continue
        seen.add(key)
        match = by_sku.get(item["sku"].upper()) if item["sku"] else None
        match_type = "SKU" if match else ""
        if not match:
            candidates = by_name.get(norm_text(item["name"]), [])
            if candidates:
                match = candidates[0]
                match_type = "NOMBRE / SKU DISTINTO"
                item["other_skus"] = ", ".join(x for x in [item["other_skus"], item["sku"]] if x)
                item["sku"] = match["sku"]
        old_price = match.get("price") if match else None
        item.update({
            "maestro_sku": match["sku"] if match else "",
            "maestro_price": old_price or "",
            "envio_grande": envio_flag(match.get("envio") if match else None),
            "match": match_type or "NO ENCONTRADO",
            "diff_pct": round((item["compra"] - old_price) / old_price * 100, 2) if numeric(old_price) else "",
            "categoria": match.get("categoria", "") if match else "",
            "inventario": match.get("inventario") if match else None,
            "canales": match.get("canales", "") if match else "",
        })
        result.append(item)
    return result


def write_workbooks(rows, out_dir, label):
    out_dir.mkdir(parents=True, exist_ok=True)
    control_path = out_dir / f"Control_{label}.xlsx"
    sync_path = out_dir / f"Nuevos_{label}_Sync.xlsx"

    control = openpyxl.Workbook()
    ws = control.active
    ws.title = "Control"
    headers = ["Fecha", "Hora", "SKU", "Otros SKUs", "Artículo", "Imagen", "Precio fuente", "Fuente", "Compra Mayorista", "Venta Bruta", "Envío Grande", "Maestro SKU", "Maestro F", "Diff %", "Coincidencia"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="34495E")
    for item in rows:
        row_num = ws.max_row + 1
        ws.append([item[k] for k in ["date", "time", "sku", "other_skus", "name", "image", "source", "source_kind", "", "", "envio_grande", "maestro_sku", "maestro_price", "diff_pct", "match"]])
        ws.cell(row_num, 9).value = f"=G{row_num}*{FACTOR_COMPRA}"
        ws.cell(row_num, 10).value = f"=I{row_num}*{FACTOR_VENTA}"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{ws.max_row}"
    for col, width in enumerate([12, 8, 22, 18, 58, 35, 16, 18, 18, 16, 14, 22, 16, 10, 22], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    style_workbook(control)
    control.save(control_path)

    sync = openpyxl.Workbook()
    sws = sync.active
    sws.title = "Nuevos Sync"
    sws.append(["SKU", "Artículo", "Precio Compra Mayorista", "Fecha Actualización", "Imágenes JPG", "Precio Venta Bruto", "Envío Grande", "Categoría", "Inventario", "Canales de Venta"])
    for item in rows:
        sws.append([item["sku"], item["name"], item["compra"], item["date"], item["image"], item["venta"], item["envio_grande"],
                    item.get("categoria", ""), item.get("inventario"), item.get("canales", "")])
    sws.freeze_panes = "A2"
    sws.auto_filter.ref = f"A1:J{sws.max_row}"
    for col, width in enumerate([22, 62, 22, 18, 35, 18, 14, 28, 12, 22], 1):
        sws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    style_workbook(sync)
    sync.save(sync_path)
    return control_path, sync_path


def main():
    parser = argparse.ArgumentParser(description="Genera Control y Sync desde un chat.txt exportado")
    parser.add_argument("chat", type=Path)
    parser.add_argument("--maestro", type=Path, default=BASE / "excel" / "Productos_Maestro.xlsx")
    parser.add_argument("--images-dir", type=Path, default=None)
    parser.add_argument("--date", default=date.today().isoformat(), help="YYYY-MM-DD para el nombre de salida")
    parser.add_argument("--out-dir", type=Path, default=BASE / "excel")
    args = parser.parse_args()
    label = datetime.strptime(args.date, "%Y-%m-%d").strftime("%d_%m_%Y")
    rows = compare_and_dedupe(parse_chat(args.chat, args.images_dir), args.maestro)
    control, sync = write_workbooks(rows, args.out_dir, label)
    print(f"Productos: {len(rows)}")
    print(f"Control: {control}")
    print(f"Sync: {sync}")


if __name__ == "__main__":
    main()
