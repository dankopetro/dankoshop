#!/usr/bin/env python3
"""
procesar_chat_14_15.py - Genera Control_14_15_Agosto.xlsx

Analiza el chat de WhatsApp del 14 y 15 de agosto y arma un Excel
"inteligente" sincronizable con el maestro.

Formatos de precio soportados:
  A) Estándar (igual al 12-13):
     Pagando en efectivo 15% OFF → precio_unidad
     Pagando por transferencia → mayorista_transfer
  B) Luciana (mensajes cortos, precio en la misma línea):
     🏛️ Transferencia 10% OFF $X → precio_trans_minorista
     💰Efectivo 15% OFF $X → precio_unidad
     🏦 10%off Transferencia $X → precio_trans_minorista
     💰 15%off Efectivo $X → precio_unidad
  C) Bulk / Sahumerios:
     Precio x1 : $X → precio_unidad
     Precio x6 : $X → mayorista_efectivo
     Precio x12 : $X → mayorista_efectivo

Uso:
  python3 scripts/procesar_chat_14_15.py
"""

import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
CHAT_FILE = Path("/home/claudio/Descargas/Mayorista/Mensajes 14 y 15 agosto/chat.txt")
CHAT_PREVIO_FILE = Path("/home/claudio/Descargas/Mayorista/Mensajes  12 y 13 agosto/chat.txt")
IMAGES_DIR = Path("/home/claudio/Descargas/Mayorista/Imagenes 14  15 agosto")
MAESTRO_PATH = BASE / "excel" / "Productos_Maestro.xlsx"
BACKUP_PREVIO_PATH = BASE / "excel" / "Productos_Maestro_backup_130826_200109.xlsx"
OUT_XLSX = BASE / "excel" / "Control_14_15_Agosto.xlsx"

FACTOR = 1.072
TOLERANCIA = 0.01  # 1% -> SIN CAMBIO

UNICODE_DIGITS = {u: str(i) for i, u in enumerate("𝟎𝟏𝟐𝟑𝟒𝟓𝟔𝟕𝟖𝟗")}

EMOJIS_PRODUCTO = ["🛒", "🛞", "📱", "📺", "🫕", "🎧", "🍺", "🏠", "🪑"]

MARKETING = [
    "nuevo ingreso", "nuevos ingresos", "nuevo producto", "ultimas", "últimas",
    "super reingreso", "llegó", "llegó el", "ultima unidad", "última unidad",
    "📢ultima", "últimas 6", "últimas 4", "últimas 3", "último en stock",
    "🚨", "🔥", "¡potencia", "hace tu pedido", "en stock", "última unid",
]


def normalize_digits(s):
    return "".join(UNICODE_DIGITS.get(ch, ch) for ch in s)


# Unicode math styled letters → plain ASCII
_UNICODE_MATH = {}
for _i, _c in enumerate("ABCDEFGHIJKLMNOPQRSTUVWXYZ"):
    # Mathematical Bold (U+1D400)
    _UNICODE_MATH[chr(0x1D400 + _i)] = _c
    # Mathematical Bold Italic (U+1D468)
    _UNICODE_MATH[chr(0x1D468 + _i)] = _c
    # Mathematical Sans-Serif (U+1D5A0)
    _UNICODE_MATH[chr(0x1D5A0 + _i)] = _c
    # Mathematical Sans-Serif Bold (U+1D5D4)
    _UNICODE_MATH[chr(0x1D5D4 + _i)] = _c
    # Mathematical Sans-Serif Bold Italic (U+1D63C)
    _UNICODE_MATH[chr(0x1D63C + _i)] = _c
    # Mathematical Monospace (U+1D670)
    _UNICODE_MATH[chr(0x1D670 + _i)] = _c
for _i, _c in enumerate("abcdefghijklmnopqrstuvwxyz"):
    _UNICODE_MATH[chr(0x1D41A + _i)] = _c
    _UNICODE_MATH[chr(0x1D482 + _i)] = _c
    _UNICODE_MATH[chr(0x1D5BA + _i)] = _c
    _UNICODE_MATH[chr(0x1D5EE + _i)] = _c
    _UNICODE_MATH[chr(0x1D656 + _i)] = _c
    _UNICODE_MATH[chr(0x1D68A + _i)] = _c
for _i in range(10):
    _UNICODE_MATH[chr(0x1D7ED + _i)] = str(_i)


def normalize_unicode_text(s):
    return "".join(_UNICODE_MATH.get(ch, ch) for ch in s)


def parse_timestamp(ts_str):
    try:
        return datetime.strptime(ts_str.strip(), "%m/%d/%y, %I:%M:%S %p")
    except ValueError:
        return None


def ts_to_folder(dt):
    return dt.strftime("%Y-%m-%d at %H.%M.%S")


def split_by_timestamp(lines):
    messages = []
    cur = None
    for line in lines:
        m = re.match(r"\[(\d+/\d+/\d+,\s+\d+:\d+:\d+\s+[AP]M)\]", line)
        if m:
            if cur is not None:
                messages.append(cur)
            ts = parse_timestamp(m.group(1))
            cur = {"timestamp": ts, "lines": [line]}
        elif cur is not None:
            cur["lines"].append(line)
    if cur is not None:
        messages.append(cur)
    return messages


def strip_timestamp(line):
    line = re.sub(r"^\[\d+/\d+/\d+,\s+\d+:\d+:\d+\s+[AP]M\]\s*\w.*?:\s*", "", line)
    line = re.sub(r"<imagen omitida>\s*", "", line)
    line = re.sub(r"<video omitido>\s*", "", line)
    line = re.sub(r"\[Reenviado[^\]]*\]\s*", "", line)
    return line


def es_marketing(line):
    low = line.lower()
    return any(m in low for m in MARKETING)


def clean_title(line):
    t = normalize_unicode_text(strip_timestamp(line))
    t = re.sub(r"^(?:[\d]{3,5})[)]\s*", "", t)
    t = re.sub(r"^(?:[\d]{3,5})[)]?\s*", "", t)
    t = re.sub(r"[\[(]N[°º]\s*[\d𝟎-𝟗]{3,5}[)\]]\s*", "", t)
    t = re.sub(r"[\[(]\s*[\d𝟎-𝟗]{3,5}[)\]]\s*", "", t)
    t = re.sub(r"[^\w\sáéíóúñÁÉÍÓÚÜü.°\"'-]", " ", t)
    t = re.sub(r"\s+", " ", t).strip(" -–")
    return t


def find_title(body_lines):
    for l in body_lines:
        t = l.strip()
        if not t:
            continue
        if any(e in t for e in EMOJIS_PRODUCTO) and not es_marketing(t):
            return clean_title(t)
    for l in body_lines:
        t = l.strip()
        if not t:
            continue
        if re.search(r"N[°º]\s*[\d𝟎-𝟗]{3,5}|\([\d𝟎-𝟗]{3,5}\)|\[[\d𝟎-𝟗]{3,5}\]", normalize_digits(t)):
            if len(t) > 8 and not es_marketing(t):
                return clean_title(t)
    for l in body_lines:
        t = l.strip()
        if t and not es_marketing(t):
            return clean_title(t)
    return ""


def extract_price_value(text):
    m = re.search(r"[$]\s*([\d]{1,3}(?:[.,][\d]{3})*(?:[.,][\d]{2})?)", text)
    if not m:
        m = re.search(r"(\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{2})?)", text)
        if not m:
            return None
    return int(m.group(1).replace(".", "").replace(",", ""))


def extract_skus(lines):
    skus = []
    for line in lines:
        ln = normalize_digits(line)
        n_match = re.findall(r"[N𝐍][°º]\s*([\d𝟎-𝟗]{3,5}(?:\s*/\s*[\d𝟎-𝟗]{3,5})*)", ln)
        if n_match:
            for g in n_match:
                skus.extend([s.strip() for s in g.split("/") if s.strip()])
            continue
        br_match = re.findall(r"\[([\d𝟎-𝟗]{3,5}(?:\s*/\s*[\d𝟎-𝟗]{3,5})*)\]", ln)
        for g in br_match:
            skus.extend([s.strip() for s in g.split("/") if s.strip()])
        pa_match = re.findall(r"\(([\d𝟎-𝟗]{3,5}(?:\s*/\s*[\d𝟎-𝟗]{3,5})*)\)", ln)
        for g in pa_match:
            skus.extend([s.strip() for s in g.split("/") if s.strip()])
        ha_match = re.findall(r"#([\d𝟎-𝟗]{3,5}(?:\s*/\s*[\d𝟎-𝟗]{3,5})*)", ln)
        for g in ha_match:
            skus.extend([s.strip() for s in g.split("/") if s.strip()])
    skus = list(dict.fromkeys(skus))
    return (skus[0] if skus else None), skus


def parse_prices(lines):
    precio_lista = None
    precio_unidad = None
    mayorista_efectivo = None
    mayorista_transfer = None
    precio_trans_minorista = None

    en_transferencia = False
    en_efectivo = False

    for l in lines:
        low = l.lower()
        v = extract_price_value(l)

        # --- Format B: Luciana's short messages (precio en la misma línea) ---
        # 🏛️ Transferencia 10% off $X / 🏛️*Transferencia 10% off $X*
        if re.search(r"transferencia\s+10%\s*off", low):
            if v is not None:
                precio_trans_minorista = v
            continue
        # 🏦 10%off Transferencia $X
        if re.search(r"10%\s*off\s+transferencia", low):
            if v is not None:
                precio_trans_minorista = v
            continue

        # 💰Efectivo 15% off $X / 💰 Efectivo 15% off $X / 💰 15%off Efectivo $X
        if re.search(r"efectivo\s+15%\s*off", low) or re.search(r"15%\s*off\s+efectivo", low):
            if v is not None and "$" in l:
                if precio_unidad is None:
                    precio_unidad = v
            continue

        # --- Format A: standard forwarded messages ---
        # "Pagando por transferencia"
        if "por transferencia" in low and ("pagando" in low or "10%" in low or "🏦" in l):
            en_transferencia = True
            en_efectivo = False
            continue
        # "Pagando en efectivo"
        if ("pagando en efectivo" in low or "pagando efectivo" in low):
            en_transferencia = False
            en_efectivo = True
            if v is not None and "$" in l and "mayorista" not in low:
                if precio_unidad is None:
                    precio_unidad = v
            continue

        # --- Format C: Bulk / Sahumerios ---
        # Precio x1 / Precio x6 / Precio x12
        m_bulk = re.search(r"precio\s+x(\d+)\s*:", low)
        if m_bulk:
            qty = int(m_bulk.group(1))
            if v is not None:
                if qty == 1:
                    if precio_unidad is None:
                        precio_unidad = v
                else:
                    if mayorista_efectivo is None:
                        mayorista_efectivo = v
            continue

        # Precio Lista / Precio lista
        if "precio lista" in low or "precio list" in low:
            if v is not None:
                precio_lista = v
            continue

        # Mayorista por transferencia
        if "mayorista por transferencia" in low:
            if v is not None:
                mayorista_transfer = v
            continue
        if en_transferencia and "mayorista" in low:
            if v is not None and "por unidad" not in low:
                mayorista_transfer = v
            continue
        if en_transferencia and ("mayorista 3" in low or "mayorista $" in low or "📦" in l):
            if v is not None and "por unidad" not in low and "efectivo" not in low:
                mayorista_transfer = v
            continue

        # Mayorista efectivo
        if "mayorista" in low and (en_efectivo or "llevando 3" in low):
            if v is not None and "por unidad" not in low and "transferencia" not in low:
                if mayorista_efectivo is None:
                    mayorista_efectivo = v
            continue

        # Precio por unidad
        if "precio por unidad" in low:
            if en_transferencia:
                continue
            if v is not None:
                precio_unidad = v
            continue

        # 15% OFF minorista
        if "15% off minorista" in low or "15%off minorista" in low:
            if v is not None and precio_unidad is None:
                precio_unidad = v
            continue
        if ("15% off" in low or "15%off" in low) and "$" in l and "mayorista" not in low:
            if v is not None and precio_unidad is None:
                precio_unidad = v
            continue

        # Minorista $ en sección efectivo
        if "minorista $" in low and en_efectivo:
            if v is not None and precio_unidad is None:
                precio_unidad = v
            continue

        # 10% off transferencia (minorista)
        if "10% off transferencia" in low or "10%off transferencia" in low:
            if v is not None:
                precio_trans_minorista = v
            continue

        # formato 1️⃣ X $x (auriculares)
        if "1️⃣" in l or re.search(r"1\s*X", l):
            if v is not None and precio_unidad is None:
                precio_unidad = v
            continue
        if "6️⃣" in l and "transferencia" in low:
            if v is not None and mayorista_transfer is None:
                mayorista_transfer = v
            continue
        if "6️⃣" in l and "efectivo" in low:
            if v is not None and mayorista_efectivo is None:
                mayorista_efectivo = v
            continue

    return {
        "precio_lista": precio_lista,
        "precio_unidad": precio_unidad,
        "mayorista_efectivo": mayorista_efectivo,
        "mayorista_transfer": mayorista_transfer,
        "precio_trans_minorista": precio_trans_minorista,
    }


def parse_message(msg):
    lines = msg["lines"]
    if not any("<imagen omitida>" in l or "<video omitido>" in l for l in lines):
        return None

    body_lines = [strip_timestamp(lines[0])] + [l.strip() for l in lines[1:] if l.strip()]

    main_sku, skus = extract_skus(lines)
    prices = parse_prices(lines)
    title = find_title(body_lines)

    return {
        "fecha": msg["timestamp"].strftime("%d/%m/%Y") if msg["timestamp"] else "",
        "hora": msg["timestamp"].strftime("%H:%M") if msg["timestamp"] else "",
        "folder": ts_to_folder(msg["timestamp"]) if msg["timestamp"] else None,
        "title": title,
        "sku": main_sku,
        "skus": skus,
        **prices,
    }


def scan_image_groups():
    groups = defaultdict(list)
    for p in IMAGES_DIR.iterdir():
        if p.suffix.lower() not in [".jpeg", ".jpg", ".png", ".webp"]:
            continue
        m = re.search(r"(\d{4}-\d{2}-\d{2} at \d{2}\.\d{2}\.\d{2})", p.name)
        if not m:
            continue
        suf = re.search(r"\((\d+)\)", p.name)
        groups[m.group(1)].append((int(suf.group(1)) if suf else 0, p.name))
    out = {}
    for k, v in groups.items():
        v.sort(key=lambda x: x[0])
        out[k] = v
    return out


def load_maestro():
    wb = openpyxl.load_workbook(MAESTRO_PATH, data_only=True)
    ws = wb.active
    d = {}
    for r in range(2, ws.max_row + 1):
        sku = ws.cell(r, 1).value
        if sku:
            d[str(sku).strip()] = (ws.cell(r, 2).value, ws.cell(r, 6).value, ws.cell(r, 9).value)
    return d


def load_maestro_names():
    """Devuelve lista de nombres normalizados del maestro para match por nombre."""
    wb = openpyxl.load_workbook(MAESTRO_PATH, data_only=True)
    ws = wb.active
    names = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if name:
            names.append(name.strip().lower())
    return names


def load_precio_anterior():
    d = {}
    if CHAT_PREVIO_FILE.exists():
        with open(CHAT_PREVIO_FILE, encoding="utf-8") as f:
            lines = f.readlines()
        messages = split_by_timestamp(lines)
        for msg in messages:
            p = parse_message(msg)
            if p is None or p["sku"] is None:
                continue
            d[p["sku"]] = {
                "unidad": p["precio_unidad"],
                "mEfectivo": p["mayorista_efectivo"],
                "mTransf": p["mayorista_transfer"],
                "lista": p["precio_lista"],
                "fecha": p["fecha"],
            }
    if BACKUP_PREVIO_PATH.exists():
        wb = openpyxl.load_workbook(BACKUP_PREVIO_PATH, data_only=True)
        ws = wb.active
        for r in range(2, ws.max_row + 1):
            sku = ws.cell(r, 1).value
            if not sku:
                continue
            s = str(sku).strip()
            if s in d:
                continue
            d[s] = {
                "unidad": ws.cell(r, 5).value,
                "mEfectivo": None,
                "mTransf": ws.cell(r, 7).value,
                "lista": ws.cell(r, 4).value,
                "fecha": "backup 13/08",
            }
    return d


def is_num(v):
    return isinstance(v, (int, float)) and v


def base_price(p, maestro_sku):
    cands = []
    if is_num(p["mayorista_transfer"]):
        cands.append(("transfer", p["mayorista_transfer"]))
    if is_num(p["precio_unidad"]):
        cands.append(("unidad", p["precio_unidad"]))
    if is_num(p["precio_lista"]):
        cands.append(("lista", p["precio_lista"]))
    if not cands:
        return None, None
    if maestro_sku and is_num(maestro_sku[2]):
        target = maestro_sku[2]
        cands.sort(key=lambda c: abs(c[1] - target))
    else:
        priority = {"transfer": 0, "unidad": 1, "lista": 2}
        cands.sort(key=lambda c: priority[c[0]])
    return cands[0][1], cands[0][0]


def estado_for(p, maestro, previo):
    sku = p["sku"]
    if sku is None:
        return "SIN SKU"
    m = maestro.get(sku)
    base, fuente = base_price(p, m)

    prev = previo.get(sku) if previo else None
    if prev is not None:
        field_map = {"transfer": "mTransf", "unidad": "unidad", "lista": "lista"}
        mejor = None
        for f in ("transfer", "unidad", "lista"):
            pv = prev.get(field_map[f])
            cur = p.get({"transfer": "mayorista_transfer",
                         "unidad": "precio_unidad",
                         "lista": "precio_lista"}[f])
            if is_num(pv) and is_num(cur):
                d = abs(cur - pv) / pv
                if mejor is None or d < mejor[0]:
                    mejor = (d, f, cur, pv)
        if mejor is not None and mejor[0] < 0.25:
            _, fuente, base, ref = mejor
            p["chat_F"] = round(base * FACTOR)
            p["fuente_base"] = fuente
            p["diff_pct"] = (base - ref) / ref * 100
            p["maestro_F"] = m[1] if m else None
            p["maestro_I"] = m[2] if m else None
            p["referencia"] = f"publicacion anterior ({prev.get('fecha','')})"
            if base > ref * (1 + TOLERANCIA):
                return "AUMENTO"
            if base < ref * (1 - TOLERANCIA):
                return "BAJA"
            return "SIN CAMBIO"

    if base is None:
        if m is None:
            return "NUEVO"
        return "SIN MAYORISTA"
    chat_f = round(base * FACTOR)
    p["chat_F"] = chat_f
    p["fuente_base"] = fuente
    if m is None:
        return "NUEVO"
    target = m[2] if is_num(m[2]) else m[1]
    if not is_num(target):
        return "SIN PRECIO MAESTRO"
    diff = (base - target) / target
    p["diff_pct"] = diff * 100
    p["maestro_F"] = m[1]
    p["maestro_I"] = m[2]
    if diff > TOLERANCIA:
        return "AUMENTO"
    if diff < -TOLERANCIA:
        return "BAJA"
    return "SIN CAMBIO"


def main():
    with open(CHAT_FILE, encoding="utf-8") as f:
        lines = f.readlines()
    messages = split_by_timestamp(lines)

    img_groups = scan_image_groups()
    used_count = defaultdict(int)
    missing_imgs = 0

    parsed = []
    for msg in messages:
        p = parse_message(msg)
        if p is None:
            continue
        imgs = []
        if p["folder"] in img_groups:
            available = img_groups[p["folder"]]
            while used_count[p["folder"]] < len(available):
                _, fname = available[used_count[p["folder"]]]
                imgs.append(fname)
                used_count[p["folder"]] += 1
                break
        if not imgs:
            missing_imgs += 1
        p["imagen"] = ", ".join(imgs)
        parsed.append(p)

    maestro = load_maestro()
    previo = load_precio_anterior()

    for p in parsed:
        p["estado"] = estado_for(p, maestro, previo)

    counts = defaultdict(int)
    for p in parsed:
        counts[p["estado"]] += 1
    print(f"Mensajes con imagen: {len(parsed)} | sin imagen: {missing_imgs}")
    for k in ["NUEVO", "AUMENTO", "BAJA", "SIN CAMBIO", "SIN MAYORISTA", "SIN SKU"]:
        if counts[k]:
            print(f"  {k}: {counts[k]}")

    # --- Filtrado ---
    maestro_names = load_maestro_names()

    # 1) Eliminar duplicados: si un SKU aparece varias veces, quedarse con la fila que tiene imagen
    seen_skus = {}
    for p in parsed:
        key = p["sku"] or p["title"]  # SIN SKU -> key por título
        if key in seen_skus:
            prev = seen_skus[key]
            if p["imagen"] and not prev["imagen"]:
                seen_skus[key] = p
        else:
            seen_skus[key] = p
    parsed = list(seen_skus.values())

    # 2) Eliminar SIN CAMBIO
    parsed = [p for p in parsed if p["estado"] != "SIN CAMBIO"]

    # 3) Eliminar SIN SKU que matcheen con maestro por nombre
    def name_matches_maestro(title):
        t = title.lower()
        for mn in maestro_names:
            if t in mn or mn in t:
                return True
        return False

    parsed = [p for p in parsed if not (p["estado"] == "SIN SKU" and name_matches_maestro(p["title"]))]

    # 4) Eliminar Tupper Gemplast (sin precio)
    parsed = [p for p in parsed if not (p["title"] and "tupper" in p["title"].lower() and "gemplast" in p["title"].lower())]

    # Re-contar después del filtrado
    counts_f = defaultdict(int)
    for p in parsed:
        counts_f[p["estado"]] += 1
    print(f"\nDespués del filtrado:")
    print(f"  Total: {len(parsed)}")
    for k in ["NUEVO", "AUMENTO", "BAJA", "SIN MAYORISTA", "SIN SKU"]:
        if counts_f[k]:
            print(f"  {k}: {counts_f[k]}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control"
    headers = [
        "Fecha", "Hora", "SKU", "Otros SKUs", "Artículo", "Imagen",
        "Precio unidad (efectivo)", "Mayorista efectivo",
        "Mayorista transferencia", "Precio Lista",
        "Base (x1.072)", "Maestro F", "Diff %", "Estado", "Referencia",
    ]
    ws.append(headers)
    from openpyxl.styles import Font, PatternFill

    hf = Font(bold=True, color="FFFFFF")
    fills = {
        "AUMENTO": "C0392B",
        "BAJA": "2874A6",
        "NUEVO": "1E8449",
        "SIN CAMBIO": "7F8C8D",
        "SIN MAYORISTA": "B7950B",
        "SIN SKU": "6C3483",
    }
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = hf
        cell.fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")

    for p in parsed:
        ws.append([
            p["fecha"], p["hora"], p["sku"],
            ", ".join(p["skus"]) if p["skus"] else "",
            p["title"], p["imagen"],
            p["precio_unidad"], p["mayorista_efectivo"],
            p["mayorista_transfer"], p["precio_lista"],
            p.get("chat_F", ""), p.get("maestro_F", ""),
            round(p["diff_pct"], 2) if "diff_pct" in p else "",
            p["estado"],
            p.get("referencia", ""),
        ])
        row = ws.max_row
        if p["estado"] in fills:
            ws.cell(row, len(headers)).fill = PatternFill(
                start_color=fills[p["estado"]], end_color=fills[p["estado"]], fill_type="solid")

    widths = [12, 8, 10, 26, 45, 40, 16, 16, 18, 14, 12, 12, 9, 14, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{ws.max_row}"

    all_used = set()
    for p in parsed:
        for fname in p["imagen"].split(", "):
            if fname:
                all_used.add(fname)
    ws2 = wb.create_sheet("Imagenes sin asignar")
    ws2.append(["Archivo"])
    for p in sorted(IMAGES_DIR.iterdir()):
        if p.name not in all_used:
            ws2.append([p.name])

    wb.save(OUT_XLSX)
    print(f"\nGuardado: {OUT_XLSX}")


if __name__ == "__main__":
    main()
