#!/usr/bin/env python3
"""
restaurar_col_i.py - Reconstruye la col I "Precio Ingresado" del maestro.

Contexto: sync_medusa_to_excel.py recreaba el Excel con openpyxl.Workbook()
(8 columnas) y borraba la col I "Precio Ingresado". Este script la reconstruye
SIN borrar nada:

- Toma el maestro ACTUAL (con todos los SKUs, incluidos los nuevos + IDs Medusa).
- Para SKUs que existen en el recalc de LibreOffice (fuente con col I original):
    -> copia el valor de col I.
- Para SKUs nuevos (no presentes en el recalc):
    -> calcula I = round(F / 1.072) para reconstruir el precio ingresado.
- Reescribe la fórmula de la col F (=IF(I{r}="","",ROUND(I{r}*1.072,0)))
  en filas 2..10000 y la col E ya queda con su fórmula.
- Guarda y luego recorre el archivo con LibreOffice headless para dejar los
  valores cacheados de F (necesarios para sync_excel_to_medusa.py).

Uso:
  python3 scripts/restaurar_col_i.py
"""

import shutil
import subprocess
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
MAESTRO_PATH = BASE / "excel" / "Productos_Maestro.xlsx"
RECALC_PATH = Path("/tmp/opencode/recalc/Productos_Maestro.xlsx")
RECALC_OUT = Path("/tmp/opencode/recalc_final/Productos_Maestro.xlsx")
FACTOR = 1.072

# Hacer backup del estado actual antes de tocar nada
backup = BASE / "excel" / "Productos_Maestro_backup_pre_coli.xlsx"
shutil.copy2(MAESTRO_PATH, backup)
print(f"Backup: {backup.name}")


def load_recall_col_i():
    """Devuelve {sku: valor_col_i} del recalc de LibreOffice."""
    if not RECALC_PATH.exists():
        print(f"WARN: no existe {RECALC_PATH}, se usarán solo cálculos F/1.072")
        return {}
    wb = openpyxl.load_workbook(RECALC_PATH, data_only=True)
    ws = wb.active
    out = {}
    for r in range(2, ws.max_row + 1):
        sku = ws.cell(r, 1).value
        if sku:
            out[str(sku).strip()] = ws.cell(r, 9).value
    return out


def main():
    wb = openpyxl.load_workbook(MAESTRO_PATH, data_only=False)
    ws = wb.active

    recalc = load_recall_col_i()
    print(f"SKUs en recalc con col I: {len(recalc)}")

    # Escribir header de col I si no existe
    if ws.cell(1, 9).value is None:
        ws.cell(1, 9).value = "Precio Ingresado"

    nuevos = []
    usados_recalc = 0
    calculados = 0
    n_skus = 0

    for r in range(2, ws.max_row + 1):
        sku = ws.cell(r, 1).value
        if not sku:
            continue
        n_skus += 1
        sku = str(sku).strip()
        f_val = ws.cell(r, 6).value

        # 1) Si está en el recalc, copiar col I
        if sku in recalc:
            ws.cell(r, 9).value = recalc[sku]
            usados_recalc += 1
        # 2) Si es nuevo y F tiene valor, derivar I = round(F/1.072)
        elif isinstance(f_val, (int, float)) and f_val:
            ws.cell(r, 9).value = int(round(float(f_val) / FACTOR))
            nuevos.append(sku)
            calculados += 1

        # 3) Reescribir fórmula de F para mantenerla viva
        ws.cell(r, 6).value = f'=IF(I{r}="","",ROUND(I{r}*1.072,0))'

    # Asegurar fórmula F en TODAS las filas hasta 10000 (incluye futuras)
    for r in range(2, 10001):
        c = ws.cell(r, 6)
        if not c.value:
            c.value = f'=IF(I{r}="","",ROUND(I{r}*1.072,0))'

    wb.save(MAESTRO_PATH)
    print(f"SKUs procesados: {n_skus}")
    print(f"  Col I copiada desde recalc: {usados_recalc}")
    print(f"  Col I derivada (nuevos): {calculados}")
    print(f"  Nuevos con I derivada ({len(nuevos)}): {sorted(nuevos)[:60]}")

    # Recalc con LibreOffice headless para cachear valores de F
    RECALC_OUT.parent.mkdir(parents=True, exist_ok=True)
    res = subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(RECALC_OUT.parent), str(MAESTRO_PATH)],
        capture_output=True, text=True, timeout=180,
    )
    if res.returncode != 0 or not RECALC_OUT.exists():
        print("ERROR LibreOffice recalc:")
        print(res.stdout[-500:])
        print(res.stderr[-500:])
        return

    shutil.copy2(RECALC_OUT, MAESTRO_PATH)
    print(f"Maestro recalculado y guardado en {MAESTRO_PATH.name}")


if __name__ == "__main__":
    main()