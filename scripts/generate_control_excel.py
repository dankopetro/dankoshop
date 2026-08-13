import openpyxl
import re
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load Maestro
wb_m = openpyxl.load_workbook('/home/claudio/Descargas/dankoshop/excel/Productos_Maestro.xlsx', data_only=True)
ws_m = wb_m.active
maestro_data = {}
for row in ws_m.iter_rows(min_row=2, values_only=True):
    sku = row[0]
    if sku:
        p_may_maestro = row[5] # Col 6: Precio Mayorista
        maestro_data[str(sku).strip()] = {
            'articulo': row[1],
            'categoria': row[2],
            'descripcion': row[3],
            'precio_lista': row[4],
            'precio_mayorista': p_may_maestro,
            'precio_maestro_calculado': p_may_maestro * 1.10 if isinstance(p_may_maestro, (int, float)) else 'N/A'
        }

# Load Luciana
wb_l = openpyxl.load_workbook('/home/claudio/Descargas/Mayorista/productos_luciana.xlsx', data_only=True)
ws_l = wb_l.active

def extract_price(row_idx):
    p_list = ws_l.cell(row_idx, 7).value
    p_may = ws_l.cell(row_idx, 6).value
    efectivo = ws_l.cell(row_idx, 8).value
    desc = str(ws_l.cell(row_idx, 5).value or '')
    
    if isinstance(p_list, (int, float)) and p_list > 1000:
        return p_list
    if isinstance(p_list, (int, float)) and 0 < p_list < 1000:
        return p_list * 1000
        
    if isinstance(efectivo, (int, float)) and efectivo > 1000:
        return efectivo
    if isinstance(p_may, (int, float)) and p_may > 1000:
        return p_may
        
    prices = re.findall(r'\$([\d\.,]+)', desc)
    if prices:
        for p_s in reversed(prices):
            clean = p_s.replace('.', '').replace(',', '')
            try:
                v = float(clean)
                if v > 1000:
                    return v
            except:
                pass
                
    if isinstance(p_list, (int, float)):
        return p_list * 1000 if p_list < 1000 else p_list
    if isinstance(efectivo, (int, float)):
        return efectivo
    if isinstance(p_may, (int, float)):
        return p_may
    return 0

updated_list = []
for r in range(2, ws_l.max_row + 1):
    sku = ws_l.cell(r, 4).value
    art = ws_l.cell(r, 3).value
    desc = ws_l.cell(r, 5).value
    p_may = ws_l.cell(r, 6).value
    efectivo = ws_l.cell(r, 8).value
    transf = ws_l.cell(r, 9).value
    ts_str = ws_l.cell(r, 10).value
    
    if not sku or not ts_str:
        continue
    
    sku_str = str(sku).strip()
    try:
        dt = datetime.strptime(str(ts_str).strip(), '%m/%d/%y %I:%M:%S %p')
    except:
        try:
            dt = datetime.strptime(str(ts_str).strip(), '%m/%d/%Y %I:%M:%S %p')
        except:
            continue
            
    if dt >= datetime(2026, 8, 8):
        m_item = maestro_data.get(sku_str, {})
        clean_price = extract_price(r)
        updated_list.append({
            'sku': sku_str,
            'articulo': art or m_item.get('articulo', ''),
            'descripcion': desc or m_item.get('descripcion', ''),
            'precio_lista_luciana': clean_price if clean_price > 0 else None,
            'precio_mayorista_luciana': p_may,
            'efectivo_luciana': efectivo,
            'transferencia_luciana': transf,
            'timestamp': ts_str,
            'precio_lista_maestro': m_item.get('precio_maestro_calculado', 'N/A'),
            'precio_mayorista_maestro': m_item.get('precio_mayorista', 'N/A')
        })

# Create new workbook for control
wb_out = openpyxl.Workbook()
ws_out = wb_out.active
ws_out.title = "Actualizados 8-Ago"

headers = [
    "SKU", "Artículo", "Descripción", 
    "Precio Lista (Luciana)", "Precio Mayorista (Luciana)", 
    "Efectivo (Luciana)", "Transferencia (Luciana)", 
    "Fecha Actualización", "Precio Lista (Maestro)", "Precio Mayorista (Maestro)",
    "Precio Para Web"
]
ws_out.append(headers)

for idx, item in enumerate(updated_list):
    row_idx = idx + 2
    p_val = item['precio_lista_luciana']
    web_formula = f"=D{row_idx}*1.25*1.072" if p_val is not None else 0
    
    ws_out.append([
        item['sku'],
        item['articulo'],
        item['descripcion'],
        item['precio_lista_luciana'],
        item['precio_mayorista_luciana'],
        item['efectivo_luciana'],
        item['transferencia_luciana'],
        item['timestamp'],
        item['precio_lista_maestro'],
        item['precio_mayorista_maestro'],
        web_formula
    ])

# Styling
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# Format header
for col in range(1, len(headers) + 1):
    cell = ws_out.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws_out.row_dimensions[1].height = 25

# Format data rows
for row_idx in range(2, len(updated_list) + 2):
    is_even = (row_idx % 2 == 0)
    current_fill = zebra_fill if is_even else white_fill
    ws_out.row_dimensions[row_idx].height = 20
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws_out.cell(row=row_idx, column=col_idx)
        cell.fill = current_fill
        cell.border = thin_border
        cell.font = Font(name="Calibri", size=10)
        
        # Alignments & Number formats
        if col_idx in [1, 8]:  # SKU, Date
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [4, 5, 6, 7, 9, 10, 11]:  # Prices & Web Price
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(cell.value, (int, float)) or (isinstance(cell.value, str) and cell.value.startswith('=')):
                cell.number_format = '$#,##0'
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Auto-adjust column widths
for col in ws_out.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if col[0].column == 3:
            max_len = 35
            break
        val_str = str(cell.value or '')
        if len(val_str) > max_len:
            max_len = len(val_str)
    ws_out.column_dimensions[col_letter].width = max(min(max_len + 4, 50), 12)

out_path = '/home/claudio/Descargas/dankoshop/excel/Control_Precios_Luciana_Desde_8_Agosto.xlsx'
wb_out.save(out_path)
print(f'Successfully updated and styled {out_path} with Column I as Maestro Column F * 1.10 and {len(updated_list)} records.')
