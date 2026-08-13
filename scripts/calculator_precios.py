#!/usr/bin/env python3
"""Calculadora de precios - Privado"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Calculadora"

input_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
result_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

ws.merge_cells("A1:C1")
ws["A1"] = "CALCULADORA DE PRECIOS"
ws["A1"].font = Font(bold=True, size=14, color="1E40AF")

ws["A3"] = "PRECIO POR UNIDAD (mayorista, ej: $4.839):"
ws["A3"].font = Font(bold=True, size=11)
ws["B3"] = 4839
ws["B3"].fill = input_fill
ws["B3"].font = Font(bold=True, size=12)
ws["B3"].number_format = '"$"#,##0'

ws["A4"] = "DESCUENTO % (ej: 17):"
ws["A4"].font = Font(bold=True, size=11)
ws["B4"] = 17
ws["B4"].fill = input_fill
ws["B4"].font = Font(bold=True, size=12)

ws["A6"] = "RESULTADOS"
ws["A6"].font = Font(bold=True, size=12, color="FFFFFF")
ws["A6"].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
ws["B6"].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
ws["C6"].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")

results = [
    ("Precio original (sin desc)", "=B3/(1-B4/100)", "Precio real si compras 1 unidad"),
    ("", "", ""),
    ("x 1.10 (+10% EFECTIVO)", "=B3/(1-B4/100)*1.10", "Tu precio venta efectivo"),
    ("x 1.15 (+15% TRANSF/DEB)", "=B3/(1-B4/100)*1.15", "Tu precio venta transferencia"),
    ("x 1.25 (+25% TARJETA)", "=B3/(1-B4/100)*1.25", "Tu precio venta tarjeta"),
]

for i, (desc, formula, note) in enumerate(results, 7):
    ws.cell(row=i, column=1, value=desc).border = border
    ws.cell(row=i, column=1).font = Font(bold=True, size=11)
    ws.cell(row=i, column=2, value=formula).border = border
    ws.cell(row=i, column=2).number_format = '"$"#,##0.00'
    ws.cell(row=i, column=2).font = Font(bold=True, size=12)
    ws.cell(row=i, column=3, value=note).border = border
    ws.cell(row=i, column=3).font = Font(color="6B7280", size=10)
    if "EFECTIVO" in desc:
        ws.cell(row=i, column=1).font = Font(bold=True, size=11, color="059669")
        ws.cell(row=i, column=2).fill = result_fill
    elif "TRANSF" in desc:
        ws.cell(row=i, column=1).font = Font(bold=True, size=11, color="2563EB")
        ws.cell(row=i, column=2).fill = result_fill
    elif "TARJETA" in desc:
        ws.cell(row=i, column=1).font = Font(bold=True, size=11, color="DC2626")
        ws.cell(row=i, column=2).fill = result_fill

ws.column_dimensions["A"].width = 40
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 38

wb.save("/home/claudio/Descargas/dankoshop/scripts/calculator_precios.xlsx")
print("Listo")
