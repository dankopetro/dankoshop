#!/usr/bin/env python3
"""
preparar_nuevos_sync.py - Genera el Excel de sync para productos nuevos.

Lee excel/Nuevos_Para_Medusa.xlsx (SKU, Artículo, Precio Ingresado, Fecha) y
genera excel/Nuevos_Para_Medusa_sync.xlsx con la estructura que espera
sync_excel_to_medusa.py:

  A (0): SKU
  B (1): Artículo
  C (2): Precio Ingresado            (referencia)
  D (3): Fecha Actualización         (referencia)
  E (4): Imágenes JPG                (referencia manual, NO se sube a Medusa)
  F (5): Precio Mayorista            = round(Ingresado * 1.072)
  G (6): Envío Grande                (vacío)

Las imágenes se toman de /home/claudio/Descargas/Mayorista/productos_luciana.xlsx
(columna Imagen) mapeando por SKU (col D) y código entre paréntesis del Artículo.
Si un SKU tiene varias fotos, se listan todas separadas por coma.

Uso:
  python3 scripts/preparar_nuevos_sync.py
"""

import re
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
NUEVOS_PATH = BASE / "excel" / "Nuevos_Para_Medusa.xlsx"
SYNC_PATH = BASE / "excel" / "Nuevos_Para_Medusa_sync.xlsx"
LUCIANA_PATH = Path("/home/claudio/Descargas/Mayorista/productos_luciana.xlsx")
FACTOR = 1.072


def leer_nuevos():
    wb = openpyxl.load_workbook(NUEVOS_PATH, data_only=True)
    ws = wb.active
    nuevos = []
    for r in range(2, ws.max_row + 1):
        sku = ws.cell(r, 1).value
        if not sku:
            continue
        nuevos.append({
            "sku": str(sku).strip(),
            "articulo": str(ws.cell(r, 2).value or "").strip(),
            "precio_ingresado": ws.cell(r, 3).value,
            "fecha": ws.cell(r, 4).value,
        })
    return nuevos


def leer_imagenes_luciana():
    wb = openpyxl.load_workbook(LUCIANA_PATH, data_only=True)
    ws = wb.active
    map_img = {}
    for r in range(2, ws.max_row + 1):
        img = ws.cell(r, 2).value
        art = str(ws.cell(r, 3).value or "")
        sku_col = str(ws.cell(r, 4).value or "").strip()
        codes = set()
        if sku_col:
            codes.add(sku_col)
        for m in re.finditer(r"\([Nn°º]?\s*(\d+)\s*\)", art):
            codes.add(m.group(1))
        for m in re.finditer(r"[Nn°º]\s*(\d{3,})", art):
            codes.add(m.group(1))
        for c in codes:
            if img:
                map_img.setdefault(c, []).append(str(img).strip())
    return map_img


def main():
    nuevos = leer_nuevos()
    print(f"Productos nuevos: {len(nuevos)}")
    map_img = leer_imagenes_luciana()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nuevos Sync"
    ws.append(["SKU", "Artículo", "Precio Ingresado", "Fecha Actualización",
               "Imágenes JPG", "Precio Mayorista", "Envío Grande"])

    sin_imagen = []
    for p in nuevos:
        imgs = list(dict.fromkeys(map_img.get(p["sku"], [])))
        if not imgs:
            sin_imagen.append(p["sku"])
        precio = p["precio_ingresado"]
        try:
            mayorista = int(round(float(precio) * FACTOR))
        except (TypeError, ValueError):
            mayorista = ""
        ws.append([
            p["sku"],
            p["articulo"],
            p["precio_ingresado"],
            p["fecha"],
            ", ".join(imgs),
            mayorista,
            "",
        ])

    wb.save(SYNC_PATH)
    print(f"Guardado: {SYNC_PATH.name}")
    if sin_imagen:
        print(f"WARN sin imagen ({len(sin_imagen)}): {sin_imagen}")


if __name__ == "__main__":
    main()