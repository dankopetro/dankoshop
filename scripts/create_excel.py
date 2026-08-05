import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

with open("data/products.json", "r") as f:
    products = json.load(f)

wb = Workbook()
ws = wb.active
ws.title = "Productos"

headers = ["SKU", "Nombre", "Categoría", "Precio Lista", "Precio Efectivo", 
           "Precio Transferencia", "Precio Mayorista", "Cuotas", "Valor Cuota",
           "Descripción", "Imágenes"]

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=11)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center")

for row_idx, p in enumerate(products, 2):
    prices = p.get("prices", {})
    imgs = len(p.get("images", []))
    
    data = [
        p.get("sku", ""),
        p.get("name", ""),
        p.get("category", ""),
        prices.get("precio_lista"),
        prices.get("precio_efectivo"),
        prices.get("precio_transferencia"),
        prices.get("precio_mayorista"),
        prices.get("cuotas"),
        prices.get("cuota_valor"),
        (p.get("description", "") or "")[:200],
        f"{imgs} imagen(es)"
    ]
    
    for col, val in enumerate(data, 1):
        cell = ws.cell(row=row_idx, column=col, value=val)
        cell.border = thin_border
        if col in (4, 5, 6, 7, 9) and val is not None:
            cell.number_format = '#,##0'
        if col == 10:
            cell.alignment = Alignment(wrap_text=True)

widths = [18, 50, 20, 14, 14, 14, 14, 8, 12, 40, 15]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[chr(64 + i) if i <= 26 else 'A' + chr(64 + i - 26)].width = w

ws.auto_filter.ref = f"A1:K{len(products) + 1}"
ws.freeze_panes = "A2"

output = "data/productos_dankoshop.xlsx"
wb.save(output)
print(f"Excel creado: {output} ({len(products)} productos)")
